from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_bullet_slide(prs: Presentation, title: str, bullets: list[tuple[str, int]]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title

    text_frame = slide.shapes.placeholders[1].text_frame
    text_frame.clear()

    for idx, (text, level) in enumerate(bullets):
        p = text_frame.paragraphs[0] if idx == 0 else text_frame.add_paragraph()
        p.text = text
        p.level = level
        p.font.size = Pt(22 if level == 0 else 18)


def add_section_divider(prs: Presentation, section_title: str, section_note: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = section_title

    box = slide.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11), Inches(2.5))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = section_note
    p.font.size = Pt(28)


def build_presentation(output_path: Path) -> None:
    prs = Presentation()

    add_title_slide(
        prs,
        "Rapido Project Teaching Deck",
        "Intelligent Mobility Insights: Ride Patterns, Cancellations, Fare Forecasting",
    )

    add_bullet_slide(
        prs,
        "How To Use This Deck",
        [
            ("This deck assumes students are complete beginners.", 0),
            ("Teach one slide, then run the exact action in VS Code or terminal.", 0),
            ("Pause after every major step and ask students to verify output.", 0),
            ("Main goal: students should understand both why and how.", 0),
        ],
    )

    add_section_divider(prs, "Part 1", "Understand the problem in simple business language")

    add_bullet_slide(
        prs,
        "Project Story In Simple English",
        [
            ("Rapido gets many bookings every day.", 0),
            ("Some rides get cancelled, some are delayed, and fares are hard to predict.", 0),
            ("If we predict these early, operations can take action before problems happen.", 0),
            ("Machine learning helps us convert past trip data into predictions.", 0),
        ],
    )

    add_bullet_slide(
        prs,
        "Business Questions We Are Solving",
        [
            ("Will this booking be completed, cancelled, or incomplete?", 0),
            ("What fare should we expect before ride starts?", 0),
            ("Is this customer likely to cancel?", 0),
            ("Is this driver likely to cause delay or incomplete trip?", 0),
        ],
    )

    add_bullet_slide(
        prs,
        "4 ML Models In This Project",
        [
            ("Model 1: Ride outcome prediction (multi-class classification)", 0),
            ("Model 2: Fare prediction (regression)", 0),
            ("Model 3: Customer cancellation risk (binary classification)", 0),
            ("Model 4: Driver delay risk (binary classification)", 0),
        ],
    )

    add_section_divider(prs, "Part 2", "Know your data before touching machine learning")

    add_bullet_slide(
        prs,
        "Dataset Files And Purpose",
        [
            ("bookings: core ride transactions and booking status", 0),
            ("customers: customer profile and cancellation behavior", 0),
            ("drivers: driver performance and delay history", 0),
            ("location_demand: demand patterns by place and hour", 0),
            ("time_features: day, weekend, season and peak indicators", 0),
        ],
    )

    add_bullet_slide(
        prs,
        "Explain Important Columns To Students",
        [
            ("booking_status is the main ride outcome label.", 0),
            ("booking_value is the fare amount we predict in regression.", 0),
            ("hour_of_day, traffic_level, weather_condition influence demand and cancellation.", 0),
            ("customer_cancel_flag and driver_delay_flag are risk labels.", 0),
        ],
    )

    add_bullet_slide(
        prs,
        "Student Checklist Before Coding",
        [
            ("Can they identify target columns for each use case?", 0),
            ("Can they explain one row from each CSV in plain language?", 0),
            ("Can they list which file connects by customer_id and driver_id?", 0),
            ("Do they understand that clean data comes before model training?", 0),
        ],
    )

    add_section_divider(prs, "Part 3", "Environment setup and project execution")

    add_bullet_slide(
        prs,
        "Step 1: Open The Project",
        [
            ("Open VS Code.", 0),
            ("Open folder PROJECT3.", 0),
            ("Check folder contains Project-files, src, artifacts, data, sql.", 0),
            ("Explain to students why structured folders improve teamwork.", 0),
        ],
    )

    add_bullet_slide(
        prs,
        "Step 2: Install Required Packages",
        [
            ("Run pip install -r requirements.txt", 0),
            ("Explain that dependencies are libraries our code needs.", 0),
            ("If any error appears, read first red line and fix one issue at a time.", 0),
        ],
    )

    add_bullet_slide(
        prs,
        "Step 3: Run Full Pipeline",
        [
            ("Run python -m src.run_pipeline", 0),
            ("This runs cleaning, feature engineering, training, and SQL export.", 0),
            ("Expected output includes model files and metrics report.", 0),
        ],
    )

    add_bullet_slide(
        prs,
        "What Students Should Verify After Pipeline",
        [
            ("data/processed/rapido_processed.csv exists", 0),
            ("artifacts/models contains 4 model files", 0),
            ("artifacts/reports/metrics.json exists", 0),
            ("sql/rapido.db exists", 0),
        ],
    )

    add_section_divider(prs, "Part 4", "Teach preprocessing and feature engineering deeply")

    add_bullet_slide(
        prs,
        "Data Cleaning: What We Did",
        [
            ("Merged date and time into booking_datetime.", 0),
            ("Filled missing actual ride time using estimated ride time.", 0),
            ("Filled missing incomplete ride reason with None.", 0),
            ("Merged all source tables using keys and time features.", 0),
        ],
    )

    add_bullet_slide(
        prs,
        "Feature Engineering: Why It Matters",
        [
            ("fare_per_km helps compare cost efficiency across rides.", 0),
            ("fare_per_min captures speed and traffic pressure.", 0),
            ("rush_hour_flag marks high-congestion hours.", 0),
            ("long_distance_flag helps distinguish trip types.", 0),
            ("city_pair captures route-level behavior.", 0),
            ("driver_reliability_score and customer_loyalty_score summarize patterns.", 0),
        ],
    )

    add_bullet_slide(
        prs,
        "Teach Leakage In Beginner Language",
        [
            ("Leakage means model sees future or answer-like information.", 0),
            ("If leakage exists, scores look very high but fail in real life.", 0),
            ("Tell students: realistic model is better than perfect fake score.", 0),
            ("Always remove direct target proxies before final evaluation.", 0),
        ],
    )

    add_section_divider(prs, "Part 5", "Model training and evaluation interpretation")

    add_bullet_slide(
        prs,
        "Modeling Pipeline Architecture",
        [
            ("Split data into train and test sets.", 0),
            ("Numeric columns: median imputation.", 0),
            ("Categorical columns: most frequent imputation + one hot encoding.", 0),
            ("Train Random Forest models for baseline performance.", 0),
        ],
    )

    add_bullet_slide(
        prs,
        "How To Explain Classification Metrics",
        [
            ("Accuracy: overall correct predictions percentage.", 0),
            ("F1 score: balance between precision and recall.", 0),
            ("AUC: ranking quality of risk probabilities.", 0),
            ("Confusion matrix: where model confuses classes.", 0),
        ],
    )

    add_bullet_slide(
        prs,
        "How To Explain Regression Metrics",
        [
            ("MAE: average absolute error in fare units.", 0),
            ("RMSE: gives higher penalty to big errors.", 0),
            ("R2: how much variance model explains.", 0),
            ("Also compare prediction error percentage against actual fare.", 0),
        ],
    )

    add_bullet_slide(
        prs,
        "Classroom Exercise: Read metrics.json",
        [
            ("Open metrics file after training.", 0),
            ("Ask students to identify best and weakest model.", 0),
            ("Ask what business action each metric supports.", 0),
            ("Discuss why extremely high metrics should be questioned.", 0),
        ],
    )

    add_section_divider(prs, "Part 6", "Visualization, dashboard, and stakeholder communication")

    add_bullet_slide(
        prs,
        "Step 4: Run Streamlit App",
        [
            ("Run streamlit run app.py", 0),
            ("Tab 1: Business KPIs and cancellation-by-hour trend.", 0),
            ("Tab 2: Single booking prediction for live demo.", 0),
            ("Tab 3: Batch CSV prediction and download output.", 0),
        ],
    )

    add_bullet_slide(
        prs,
        "How To Present Insights To Non-Technical Audience",
        [
            ("Start with business problem, not algorithm name.", 0),
            ("Use before and after story: what action this model enables.", 0),
            ("Show one chart, one risk score, one recommended action.", 0),
            ("End with measurable impact target (for example 20 percent cancellation reduction).", 0),
        ],
    )

    add_bullet_slide(
        prs,
        "Operational Interventions Examples",
        [
            ("High customer cancellation risk: send ETA transparency or incentive.", 0),
            ("High driver delay risk: allocate backup driver early.", 0),
            ("High demand window: apply surge and pre-position drivers.", 0),
            ("Low reliability route pair: monitor and improve matching strategy.", 0),
        ],
    )

    add_section_divider(prs, "Part 7", "Project submission and student evaluation readiness")

    add_bullet_slide(
        prs,
        "Final Submission Checklist",
        [
            ("Code runs end to end without manual edits.", 0),
            ("README explains setup and run commands clearly.", 0),
            ("Metrics and charts are included in final report.", 0),
            ("Students can explain one business decision per model.", 0),
        ],
    )

    add_bullet_slide(
        prs,
        "Common Student Errors And Fixes",
        [
            ("Error: module not found. Fix: install requirements in active environment.", 0),
            ("Error: file path issue. Fix: run commands from project root folder.", 0),
            ("Error: unrealistic perfect scores. Fix: review leakage-prone features.", 0),
            ("Error: presentation confusion. Fix: explain use case first, metric later.", 0),
        ],
    )

    add_bullet_slide(
        prs,
        "Suggested 2-Hour Teaching Plan",
        [
            ("0-20 min: Business context and use cases", 0),
            ("20-50 min: Data files and preprocessing walkthrough", 0),
            ("50-90 min: Training pipeline and metric interpretation", 0),
            ("90-120 min: Streamlit demo and Q and A", 0),
        ],
    )

    add_bullet_slide(
        prs,
        "Thank You Slide",
        [
            ("Students should now be able to run, explain, and present the project.", 0),
            ("Next class task: remove leakage and compare realistic model metrics.", 0),
            ("Optional extension: build FastAPI endpoint for live prediction.", 0),
        ],
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)


if __name__ == "__main__":
    target = Path("docs") / "Rapido_Student_Teaching_Deck.pptx"
    build_presentation(target)
    print(f"Created: {target.resolve()}")
